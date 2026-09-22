# -*- coding: utf-8 -*-
"""Xuất một bảng tính Excel để đưa lên Google Sheet và đồng bộ ngược về cổng.

    set PORTAL_PASS=...
    python scripts/xuat_bang_tinh_dong_bo.py                       # máy cục bộ, kỳ hiện tại
    python scripts/xuat_bang_tinh_dong_bo.py live 2026-09          # đọc số từ bản live
    python scripts/xuat_bang_tinh_dong_bo.py local 2026-09 "D:/bang.xlsx"

Mỗi đầu việc một tab, tên tab đặt đúng tên đầu việc để cổng nhận ra khi đọc
ngược; tab đã điền sẵn số đang có nên mở lên là sửa, không phải gõ lại. Cột đặt
đúng khuôn cổng đọc:

  - đầu việc theo cụm : ky, trung_tam, ke_hoach, thuc_hien, bkk, ghi_chu
  - hoàn công         : ky, nhom, trang_thai, sl_mct, cong_no, ghi_chu
  - dự án             : ky, khoi, don_vi, tong_trien_khai, <7 bước>, ghi_chu

Mỗi bảng còn có cột "dau_viec" ghi mã đầu việc ở từng dòng. Tên tab của Excel
tối đa 31 ký tự nên tên đầu việc dài bị cắt, hai đầu việc tên gần giống nhau dễ
lẫn; cổng đọc mã trong cột này trước, nên cắt tên hay đổi tên tab đều không sai
chỗ. Script vẫn nêu ra những tab mà riêng tên tab thì còn nhập nhằng.
"""
import io
import os
import re
import sys
import unicodedata
from datetime import date

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

DICH = {"local": "http://127.0.0.1:8000", "live": "https://portal-vhkt-backend.n1.tinhgon.xyz"}
NOI = sys.argv[1] if len(sys.argv) > 1 else "local"
KY = sys.argv[2] if len(sys.argv) > 2 else f"{date.today():%Y-%m}"
RA = sys.argv[3] if len(sys.argv) > 3 else f"bang-tinh-dong-bo-{KY}.xlsx"
B = DICH.get(NOI, NOI)
A = f"{B}/api/admin"

DAU = Font(bold=True, color="FFFFFF")
NEN = PatternFill("solid", fgColor="C8102E")
NEN_PHU = PatternFill("solid", fgColor="F2F2F2")


def _slug(s: str) -> str:
    s = unicodedata.normalize("NFD", (s or "").strip().lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").replace("đ", "d")
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def lay(s, duong, **t):
    r = s.get(f"{A}{duong}", params=t or None, timeout=120)
    if r.status_code >= 400:
        raise SystemExit(f"GET {duong} -> {r.status_code} {r.text[:200]}")
    return r.json()


def ten_tab(nhan: str, da_dung: set) -> str:
    """Tên tab hợp lệ cho Excel: bỏ ký tự cấm, tối đa 31 ký tự, không trùng."""
    ten = re.sub(r"[\\/*?:\[\]]", "-", nhan).strip()[:31].strip()
    goc, i = ten, 2
    while ten.lower() in da_dung:
        duoi = f" {i}"
        ten = (goc[:31 - len(duoi)] + duoi).strip()
        i += 1
    da_dung.add(ten.lower())
    return ten


def ghi_bang(ws, cot, dong):
    ws.append(cot)
    for i, _c in enumerate(cot, start=1):
        o = ws.cell(row=1, column=i)
        o.font, o.fill = DAU, NEN
        o.alignment = Alignment(horizontal="center", vertical="center")
    for d in dong:
        ws.append([d.get(c, "") for c in cot])
    for i, c in enumerate(cot, start=1):
        rong = max([len(str(c))] + [len(str(d.get(c, ""))) for d in dong] or [8])
        ws.column_dimensions[get_column_letter(i)].width = min(max(rong + 2, 10), 46)
    ws.freeze_panes = "A2"


s = requests.Session()
dn = s.post(f"{B}/api/auth/login",
            data={"username": "admin", "password": os.environ["PORTAL_PASS"]}, timeout=60).json()
if "access_token" not in dn:
    raise SystemExit(f"Không đăng nhập được {B}: {dn}")
s.headers["Authorization"] = f"Bearer {dn['access_token']}"

cats = lay(s, "/tech-tasks/categories")
buoc = [b["ma"] for b in (lay(s, "/du-an/x", period=KY).get("buoc") or [])]
print(f"== {B} · kỳ {KY} · {len(cats)} đầu việc")

wb = Workbook()
wb.remove(wb.active)
da_dung, ban_do, bo_qua = set(), [], []

# ---- trang hướng dẫn đứng đầu ----
hd = wb.create_sheet("HƯỚNG DẪN")
for d in [
    ["BẢNG TÍNH ĐỒNG BỘ SỐ LIỆU — CÔNG VIỆC MẢNG KỸ THUẬT"],
    [f"Kỳ báo cáo: {KY}    ·    Xuất lúc: {date.today():%d/%m/%Y}    ·    Nguồn: {B}"],
    [],
    ["Cách dùng"],
    ["1. Tải tệp này lên Google Drive rồi mở bằng Google Sheets (Tệp > Nhập > Tải lên)."],
    ["2. Chia sẻ: Bất kỳ ai có đường liên kết — vai trò Người xem."],
    ["3. Trên cổng, mở một đầu việc > Google Sheet > dán link > Kiểm tra kết nối."],
    ["4. Gõ tên các tab cần đọc vào ô “Tên tab” (mỗi tab một dòng) rồi bấm Đọc mọi tab;"],
    ["   hoặc bật “Tự động đọc” để máy chủ tự đọc lại theo chu kỳ."],
    [],
    ["Quy tắc"],
    ["· Mỗi tab mang tên một đầu việc; sửa số trong tab nào thì vào đúng đầu việc đó."],
    ["· Cột dau_viec là mã đầu việc — ĐỪNG sửa, đừng xóa: cổng dựa vào đó để đọc đúng chỗ."],
    ["· KHÔNG đổi tên tab và KHÔNG đổi dòng tiêu đề — cổng nhận cột theo tên."],
    ["· Cột ky để trống thì cổng hiểu là kỳ đang xem; điền sẵn để chắc chắn."],
    ["· Thêm dòng mới cứ thêm ở cuối; xóa dòng thì số cũ trên cổng vẫn còn."],
    ["· Ô trống được hiểu là 0."],
    [],
    ["Danh sách tab"],
]:
    hd.append(d)
hd["A1"].font = Font(bold=True, size=14, color="C8102E")
hd.append(["Tab", "Đầu việc", "Kiểu theo dõi", "Số dòng"])
for i in range(1, 5):
    o = hd.cell(row=hd.max_row, column=i)
    o.font, o.fill = DAU, NEN

KIEU_NHAN = {"cum": "Theo cụm/trung tâm", "hoan_cong": "Hoàn công", "du_an": "Dự án"}

for c in cats:
    ma, nhan, kieu = c["id"], c["label"], (c.get("kieu") or "cum")
    if kieu == "hoan_cong":
        d = lay(s, f"/hoan-cong/{ma}", period=KY)
        cot = ["dau_viec", "ky", "nhom", "trang_thai", "sl_mct", "cong_no", "ghi_chu"]
        dong = [{"dau_viec": ma, "ky": KY, "nhom": g["nhom"], "trang_thai": t["nhan"],
                 "sl_mct": t.get("sl_mct") or 0, "cong_no": t.get("cong_no") or 0,
                 "ghi_chu": t.get("note") or ""}
                for g in d.get("nhom", []) for t in g.get("trang_thai", [])]
    elif kieu == "du_an":
        d = lay(s, f"/du-an/{ma}", period=KY)
        cot = ["dau_viec", "ky", "khoi", "don_vi", "tong_trien_khai"] + buoc + ["ghi_chu"]
        dong = [{**{"dau_viec": ma, "ky": KY, "khoi": r.get("khoi") or "", "don_vi": r["don_vi"],
                    "tong_trien_khai": r.get("tong_trien_khai") or 0,
                    "ghi_chu": r.get("note") or ""},
                 **{b: r.get(b) or 0 for b in buoc}}
                for r in d.get("dong", [])]
    else:
        hm = lay(s, "/progress/items", category=ma)
        cot = ["dau_viec", "ky", "trung_tam", "ke_hoach", "thuc_hien", "bkk", "ghi_chu"]
        dong = []
        for h in hm:
            for r in lay(s, "/progress/centers", category=ma, item=h["id"], period=KY):
                if r.get("ft_name"):
                    continue
                dong.append({"dau_viec": ma, "ky": KY, "trung_tam": r["center"],
                             "ke_hoach": r.get("plan_qty") or 0, "thuc_hien": r.get("done_qty") or 0,
                             "bkk": r.get("bkk_qty") or 0, "ghi_chu": r.get("note") or ""})

    if not dong:
        # Tab chưa có số: vẫn để lại một dòng mang mã đầu việc (số để trống) để
        # cổng biết tab này thuộc về ai, và người dùng có chỗ gõ vào.
        dong = [{c: (ma if c == "dau_viec" else KY if c == "ky" else "") for c in cot}]
    ten = ten_tab(nhan, da_dung)
    ghi_bang(wb.create_sheet(ten), cot, dong)
    ban_do.append((ten, nhan, KIEU_NHAN.get(kieu, kieu), len(dong)))
    hd.append([ten, nhan, KIEU_NHAN.get(kieu, kieu), len(dong)])
    if _slug(ten) != _slug(nhan):
        bo_qua.append((ten, nhan))

for i, r in enumerate([22, 46, 20, 10], start=1):
    hd.column_dimensions[get_column_letter(i)].width = r

# ---- tên tab bị cắt có còn nhận ra đúng một đầu việc không ----
canh_bao = []
for ten, nhan in bo_qua:
    goc = _slug(ten)
    khop = [c["label"] for c in cats
            if _slug(c["label"]) == goc or goc in _slug(c["label"]) or _slug(c["label"]) in goc]
    if len(khop) != 1:
        canh_bao.append(f"Tab “{ten}” (của “{nhan}”) hợp với {len(khop)} đầu việc: {', '.join(khop[:4])}")

wb.save(RA)
print(f"Đã ghi {RA} — {len(ban_do)} tab")
print(f"  tab bị cắt tên (quá 31 ký tự): {len(bo_qua)}")
for x in canh_bao:
    print("  CẢNH BÁO:", x)
print("KẾT LUẬN: mọi tab đều có cột “dau_viec” mang mã đầu việc nên đọc về đúng chỗ"
      + ("" if not canh_bao else
         f"; riêng {len(canh_bao)} tab thì tên tab không đủ phân biệt — cứ giữ cột mã là được"))
