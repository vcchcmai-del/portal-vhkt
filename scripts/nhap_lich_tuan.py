"""Nhập lịch công tác tuần của chi nhánh (tệp "CNCTHCM- Lịch tuần (...)") lên portal.

    python scripts/nhap_lich_tuan.py "<tệp .csv hoặc .xlsx>"          # chạy thử, không ghi
    python scripts/nhap_lich_tuan.py "<tệp .csv hoặc .xlsx>" commit   # ghi thật

Ghi lên cả hai nơi: máy đang làm (L) và bản chạy thật (W). Mật khẩu admin lấy
từ biến môi trường PORTAL_PASS.

Khuôn tệp gốc (mỗi tuần chi nhánh gửi đúng mẫu này):
  - vài dòng tiêu đề/phê duyệt ở đầu, rồi hai dòng tiêu đề cột:
      Thời gian | (giờ) | Địa điểm họp | Chủ đề | Thành phần ...    | Mã CTH | Ghi chú
                |       |              |        | Chủ trì | Chuẩn bị | Thành phần tham gia
  - cột "Thời gian" chỉ ghi ở cuộc họp đầu mỗi ngày ("Thứ hai - 14/09/2026"),
    các dòng dưới để trống nghĩa là cùng ngày;
  - giờ dạng khoảng "07:00-09:00", "16h00-18h00", "18:00 - 20:00".
Cột được dò theo tên tiêu đề, nên mẫu có chèn/đổi chỗ cột vẫn đọc đúng.

Portal chỉ lưu giờ bắt đầu, nên giờ kết thúc, phần chuẩn bị và ghi chú được
ghép vào ô thành phần. Hình thức họp: có link thật thì là zoom/google_meet;
nhắc Zoom/online/CTH mà không có link thì là "khac" (mã CTH đưa vào ô mã họp);
còn lại là họp trực tiếp. Nhập theo kiểu upsert (khoá: giờ bắt đầu + nội dung),
nên nhập lại cùng tệp sau khi chi nhánh sửa lịch không sinh cuộc họp trùng.
"""
import csv
import io
import os
import re
import sys
from collections import Counter

import requests

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8")  # thông báo lỗi tiếng Việt

L = "http://127.0.0.1:8000"
W = "https://portal-vhkt-backend.n1.tinhgon.xyz"

COT_RA = ["ngay", "gio", "noi_dung", "dia_diem", "chu_tri", "thanh_phan",
          "hinh_thuc", "thong_tin_hop", "ma_hop", "mat_khau_hop"]

# tên cột trong tệp gốc (viết thường, bỏ khoảng trắng thừa) -> vai trò
TEN_COT = {
    "thời gian": "ngay", "địa điểm họp": "dia_diem", "địa điểm": "dia_diem",
    "chủ đề": "chu_de", "nội dung": "chu_de", "chủ trì": "chu_tri",
    "chuẩn bị": "chuan_bi", "thành phần tham gia": "thanh_phan",
    "mã cth": "ma_cth", "ghi chú": "ghi_chu",
}


def doc_bang(duong_dan):
    if duong_dan.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(duong_dan, data_only=True).active
        return [["" if v is None else str(v) for v in hang]
                for hang in ws.iter_rows(values_only=True)]
    with open(duong_dan, encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def tim_cot(bang):
    """Tìm dòng tiêu đề (dòng có "Chủ đề") và vị trí từng cột theo tên."""
    for i, hang in enumerate(bang):
        if any(c.strip().lower() == "chủ đề" for c in hang):
            cot = {}
            for hang_td in bang[i:i + 2]:
                for j, c in enumerate(hang_td):
                    vai = TEN_COT.get(re.sub(r"\s+", " ", c.strip().lower()))
                    if vai and vai not in cot:
                        cot[vai] = j
            # cột giờ không có tên: nằm ngay sau cột "Thời gian"
            cot["gio"] = cot.get("ngay", 0) + 1
            thieu = {"ngay", "chu_de"} - cot.keys()
            if thieu:
                sys.exit(f"Không thấy cột {thieu} ở dòng tiêu đề {i + 1}.")
            # dữ liệu bắt đầu sau dòng tiêu đề phụ (Chủ trì/Chuẩn bị...)
            bat_dau = i + 2 if "chu_tri" in cot and cot["chu_tri"] < len(bang[i + 1]) \
                and bang[i + 1][cot["chu_tri"]].strip() else i + 1
            return cot, bat_dau
    sys.exit("Không thấy dòng tiêu đề có cột \"Chủ đề\" — tệp có đúng mẫu lịch tuần không?")


def chuyen(bang):
    cot, dau = tim_cot(bang)
    ra, ngay, bo = [], None, 0

    def o(hang, vai):
        j = cot.get(vai)
        return hang[j].strip() if j is not None and j < len(hang) else ""

    for hang in bang[dau:]:
        m = re.search(r"(\d{1,2}/\d{1,2}/\d{4})", o(hang, "ngay"))
        if m:
            ngay = m.group(1)
        chu_de = o(hang, "chu_de")
        if not chu_de:
            if any(c.strip() for c in hang):
                bo += 1
            continue
        if not ngay:
            sys.exit(f"Cuộc họp \"{chu_de}\" nằm trước dòng ghi ngày đầu tiên.")
        khoang = [p.strip().replace("h", ":").replace("H", ":")
                  for p in re.split(r"\s*[-–]\s*", o(hang, "gio")) if p.strip()]
        bat_dau = khoang[0] if khoang else ""
        ket_thuc = khoang[1] if len(khoang) > 1 else ""
        dia_diem, ma_cth = o(hang, "dia_diem"), o(hang, "ma_cth")
        tat_ca = " ".join([dia_diem, o(hang, "thanh_phan"), chu_de, o(hang, "ghi_chu")])
        link = re.search(r"https?://\S+", tat_ca)
        ma_id = re.search(r"Meeting ID:\s*([\d ]+\d)", tat_ca)
        cth = re.search(r"CTH:\s*([^)]+)\)", dia_diem)
        if link:
            hinh_thuc = "google_meet" if "meet.google" in link.group(0) else "zoom"
            thong_tin = link.group(0)
        else:
            hinh_thuc = "khac" if (re.search(r"zoom|meet|online|CTH", tat_ca, re.I) or ma_cth) \
                else "truc_tiep"
            thong_tin = ""
        ma_hop = (ma_id.group(1) if ma_id else "") or (cth.group(1).strip() if cth else "") or ma_cth
        thanh_phan = "; ".join(x for x in [
            f"Kết thúc {ket_thuc}" if ket_thuc else "",
            f"Chuẩn bị: {o(hang, 'chuan_bi')}" if o(hang, "chuan_bi") else "",
            o(hang, "thanh_phan"),
            f"Ghi chú: {o(hang, 'ghi_chu')}" if o(hang, "ghi_chu") else "",
        ] if x)
        ra.append([ngay, bat_dau, chu_de, dia_diem, o(hang, "chu_tri"), thanh_phan,
                   hinh_thuc, thong_tin, ma_hop, ""])
    return ra, bo


def nhap(base, du_lieu_csv, ghi):
    mat_khau = os.environ.get("PORTAL_PASS")
    if not mat_khau:
        sys.exit("Đặt mật khẩu admin vào biến môi trường PORTAL_PASS trước khi chạy.")
    s = requests.Session()
    dn = s.post(f"{base}/api/auth/login", data={"username": "admin", "password": mat_khau},
                timeout=30)
    dn.raise_for_status()
    kq = s.post(f"{base}/api/admin/import/file",
                headers={"Authorization": f"Bearer {dn.json()['access_token']}"},
                params={"kind": "schedule", "mode": "upsert", "commit": str(ghi).lower()},
                files={"file": ("lich-tuan.csv", du_lieu_csv, "text/csv")}, timeout=120).json()
    print(f"{base} | {'GHI' if ghi else 'chạy thử'} |", kq.get("tom_tat") or kq)
    for d in kq.get("chi_tiet", []):
        if d.get("trang_thai") == "loi":
            print("   LỖI dòng", d.get("dong"), ":", d.get("thong_bao"))


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    ghi = len(sys.argv) > 2 and sys.argv[2] == "commit"
    ra, bo = chuyen(doc_bang(sys.argv[1]))
    print(f"Đọc {len(ra)} cuộc họp | theo ngày: {dict(Counter(x[0] for x in ra))}"
          f" | hình thức: {dict(Counter(x[6] for x in ra))} | dòng lẻ bỏ qua: {bo}")
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COT_RA)
    w.writerows(ra)
    du_lieu = buf.getvalue().encode("utf-8-sig")
    for base in (L, W):
        nhap(base, du_lieu, ghi)


if __name__ == "__main__":
    main()
