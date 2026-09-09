"""Tra cứu CSDL hạ tầng mạng lưới — đọc sheet "TỔNG HỢP" của file Excel báo
cáo hạ tầng hàng tháng (nhiều bảng con xếp chồng, mỗi bảng 1 nhóm chỉ tiêu
theo Mã cụm) và làm phẳng thành các dòng (kỳ, nhóm, trung tâm, chỉ tiêu, giá
trị) để tra cứu theo trung tâm hoặc theo tên chỉ tiêu.

Không dùng chung với bulkimport.py/models.Metric — xem lý do trong kế hoạch:
sheet nguồn có cấu trúc nhiều bảng con không khớp khuôn "1 tiêu đề + N dòng"
của bulkimport, và Metric bị khoá quyền theo module "dashboard" (không phải
"csdl_ht" mà module này cần).
"""
import io
import re
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import PlainTextResponse
from sqlalchemy.orm import Session

from . import models
from .auditlog import log_action
from .database import get_db
from .permissions import require_module

admin_router = APIRouter(prefix="/api/admin", tags=["CSDL hạ tầng"])

ROMAN_PREFIX_RE = re.compile(r"^(VIII|VII|VI|IV|IX|X|I|II|III|V)[\.:\s]")
PERIOD_RE = re.compile(r"THÁNG\s*0?(\d{1,2})[\-/](\d{4})", re.IGNORECASE)


def _clean_cell(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def _is_data_row_marker(a, center) -> bool:
    """Dòng dữ liệu: cột Mã cụm khác rỗng, và cột STT là số hoặc "Tổng"."""
    if center is None:
        return False
    if isinstance(a, (int, float)):
        return True
    if isinstance(a, str) and a.strip().casefold() in ("tổng", "tong"):
        return True
    return False


def _merge_map(ws):
    """{(dòng,cột) 1-based: giá trị} cho MỌI ô trong vùng gộp — openpyxl chỉ
    trả giá trị ở ô góc trên-trái vùng gộp, các ô còn lại đọc ra None. Cần
    workbook mở KHÔNG ở chế độ read_only (mới có ws.merged_cells)."""
    m = {}
    for rng in ws.merged_cells.ranges:
        top_val = ws.cell(row=rng.min_row, column=rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = top_val
    return m


def parse_tonghop(ws):
    """Trả về (period, stats, blocks) — stats là list dict sẵn sàng ghi vào
    models.InfraStat, blocks là danh sách bảng con đã nhận diện (dùng cho
    bản xem trước, để người dùng đối chiếu lại với file gốc)."""
    merges = _merge_map(ws)
    rows = list(ws.iter_rows(values_only=True))
    n = len(rows)

    def raw(i, j):
        """Giá trị THẬT của ô (dòng i, cột j — 0-based) — None nếu ô nằm
        trong vùng gộp nhưng không phải góc trên-trái. Dùng để NHẬN DIỆN
        dòng tiêu đề lớn/dòng dữ liệu, không bị nhiễu bởi vùng gộp rộng."""
        row = rows[i]
        return _clean_cell(row[j]) if j < len(row) else None

    def cell(i, j):
        """Giá trị ô đã giải quyết vùng gộp — dùng khi GHÉP nhãn cột tiêu đề."""
        v = merges.get((i + 1, j + 1))
        if v is not None:
            return _clean_cell(v)
        return raw(i, j)

    period = None
    title0 = raw(0, 0) if rows else None
    if title0:
        m = PERIOD_RE.search(title0)
        if m:
            period = f"{m.group(2)}-{int(m.group(1)):02d}"

    major_section = ""
    minor_section = ""
    stats = []
    blocks = []

    i = 0
    while i < n:
        a = raw(i, 0)
        nonnull = sum(1 for j in range(len(rows[i])) if raw(i, j) is not None)

        if isinstance(a, str) and a.strip().casefold() == "stt":
            # Một số bảng con còn có 1 dòng tiêu đề NHÓM CHA nằm NGAY TRÊN
            # dòng "STT" (vd "Thuê bao di động" phía trên "Trạm> 10,000 TB",
            # hay "5G"/"4G"/"3G"/"2G" phía trên các cột "Tổng thuê bao" —
            # thiếu dòng này thì nhiều cột trùng tên, không phân biệt được).
            # Dò ngược tối đa 2 dòng, dừng lại khi gặp dòng trống hoặc dòng
            # tiêu đề mục (đã được xử lý ở lượt trước, không gộp lại).
            pre_header_rows = []
            b = i - 1
            while b >= 0 and len(pre_header_rows) < 2:
                b_nonnull = sum(1 for j in range(len(rows[b])) if raw(b, j) is not None)
                if b_nonnull == 0 or b_nonnull <= 3:
                    break
                pre_header_rows.append(b)
                b -= 1
            pre_header_rows.reverse()

            # Gộp mọi dòng tiêu đề liên tiếp phía dưới (cột A rỗng) — sheet
            # này có bảng con tiêu đề trải 1, 2 hoặc 3 dòng khác nhau.
            header_rows = pre_header_rows + [i]
            h = i + 1
            while h < n and len(header_rows) < 4:
                ha = raw(h, 0)
                if ha is not None:
                    break
                if all(raw(h, j) is None for j in range(len(rows[h]))):
                    break
                header_rows.append(h)
                h += 1
            data_start = h

            # Chỉ nhận bảng con thật sự theo Mã cụm — vài bảng cuối sheet
            # (thống kê CCDC theo tháng/năm) dùng cột B khác hẳn ("2026",
            # "Tháng 1"...), không phải "Mã cụm", nhận nhầm sẽ lẫn dữ liệu
            # rác vào danh sách trung tâm.
            col1_header = " ".join(str(cell(r, 1)) for r in header_rows if cell(r, 1))
            if "cụm" not in col1_header.casefold():
                blocks.append({"section": f"{major_section} — {minor_section}" if minor_section else major_section,
                               "row": i + 1, "header_rows": len(header_rows), "columns": 0, "data_rows": 0,
                               "skipped": True, "skip_reason": f'Cột B không phải "Mã cụm" (là "{col1_header}").'})
                i = data_start
                continue

            max_col = max(len(rows[r]) for r in header_rows)
            labels = {}
            for j in range(2, max_col):
                parts = []
                for r in header_rows:
                    v = cell(r, j)
                    if v and (not parts or parts[-1] != v):
                        parts.append(str(v))
                if parts:
                    labels[j] = " - ".join(parts)

            section = f"{major_section} — {minor_section}" if minor_section else major_section

            k = data_start
            row_count = 0
            while k < n:
                da, dcenter = raw(k, 0), raw(k, 1)
                if not _is_data_row_marker(da, dcenter):
                    break
                center = str(dcenter)
                for j, label in labels.items():
                    val = raw(k, j)
                    if val is None:
                        continue
                    if isinstance(val, (int, float)):
                        stats.append({"period": period, "section": section, "center": center,
                                      "label": label, "value": float(val), "value_text": None})
                    else:
                        stats.append({"period": period, "section": section, "center": center,
                                      "label": label, "value": None, "value_text": str(val)})
                row_count += 1
                k += 1

            blocks.append({"section": section, "row": i + 1, "header_rows": len(header_rows),
                           "columns": len(labels), "data_rows": row_count})
            i = k
            continue

        if isinstance(a, str) and nonnull <= 3:
            text = a.strip()
            if i > 0:
                if ROMAN_PREFIX_RE.match(text):
                    major_section, minor_section = text, ""
                else:
                    minor_section = text
            i += 1
            continue

        i += 1

    return period, stats, blocks


def parse_ma_cum(ws):
    """Đọc sheet "DM Mã cụm": (Mã cụm, Tên trung tâm, Mã cũ, Tên huyện chuẩn cũ)."""
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for row in rows[2:]:
        code = _clean_cell(row[0]) if len(row) > 0 else None
        name = _clean_cell(row[1]) if len(row) > 1 else None
        old_code = _clean_cell(row[2]) if len(row) > 2 else None
        old_name = _clean_cell(row[3]) if len(row) > 3 else None
        if code and name:
            out.append((str(code), str(name), str(old_code) if old_code else None, str(old_name) if old_name else None))
    return out


@admin_router.post("/csdl-ht/import")
async def import_csdl_ht(file: UploadFile = File(...), commit: bool = False,
                         db: Session = Depends(get_db), user=Depends(require_module("csdl_ht", "create")),
                         request: Request = None):
    """Đọc file Excel CSDL hạ tầng (.xlsx). commit=false chỉ xem trước, không ghi."""
    content = await file.read(15 * 1024 * 1024 + 1)
    if len(content) > 15 * 1024 * 1024:
        raise HTTPException(400, "Tệp vượt quá 15 MB.")
    try:
        # Không dùng read_only=True: cần ws.merged_cells để giải quyết đúng
        # các ô tiêu đề gộp (rất nhiều trong sheet "TỔNG HỢP").
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"Không đọc được tệp Excel: {e}")

    if "TỔNG HỢP" not in wb.sheetnames:
        raise HTTPException(400, 'Tệp không có sheet "TỔNG HỢP".')

    period, stats, blocks = parse_tonghop(wb["TỔNG HỢP"])
    if not period:
        raise HTTPException(400, 'Không đọc được kỳ báo cáo — dòng tiêu đề đầu sheet "TỔNG HỢP" '
                                  'phải có dạng "...THÁNG 08-2026".')

    centers = parse_ma_cum(wb["DM Mã cụm"]) if "DM Mã cụm" in wb.sheetnames else []

    preview = {
        "period": period, "blocks": blocks, "total_rows": len(stats),
        "centers_found": len(centers), "sample": stats[:20],
    }
    if not commit:
        return {"committed": False, **preview}

    db.query(models.InfraStat).filter(models.InfraStat.period == period).delete()
    for s in stats:
        db.add(models.InfraStat(**s))
    for code, name, old_code, old_name in centers:
        row = db.query(models.InfraCenterCode).filter(models.InfraCenterCode.code == code).first()
        if row:
            row.name, row.old_code, row.old_name = name, old_code, old_name
        else:
            db.add(models.InfraCenterCode(code=code, name=name, old_code=old_code, old_name=old_name))
    db.commit()
    log_action(db, user, "import", "csdl_ht", None, period,
              detail=f"Ghi đè {len(stats)} dòng chỉ tiêu, kỳ {period}.", request=request)
    return {"committed": True, **preview}


@admin_router.get("/csdl-ht/periods")
def list_periods(db: Session = Depends(get_db), _=Depends(require_module("csdl_ht", "view"))):
    rows = db.query(models.InfraStat.period).distinct().order_by(models.InfraStat.period.desc()).all()
    return [r[0] for r in rows]


@admin_router.get("/csdl-ht/centers")
def list_centers(db: Session = Depends(get_db), _=Depends(require_module("csdl_ht", "view"))):
    names = {c.code: c.name for c in db.query(models.InfraCenterCode).all()}
    codes = sorted(r[0] for r in db.query(models.InfraStat.center).filter(models.InfraStat.center != "Tổng").distinct().all())
    return [{"code": c, "name": names.get(c, c)} for c in codes]


def _latest_period(db: Session) -> Optional[str]:
    row = db.query(models.InfraStat.period).order_by(models.InfraStat.period.desc()).first()
    return row[0] if row else None


@admin_router.get("/csdl-ht/sections")
def list_sections(period: Optional[str] = None, db: Session = Depends(get_db),
                  _=Depends(require_module("csdl_ht", "view"))):
    """Danh sách nhóm (I. Vô tuyến, II. Hạ tầng...) của 1 kỳ, đúng thứ tự xuất
    hiện trong file gốc (theo id nhỏ nhất của mỗi nhóm)."""
    period = period or _latest_period(db)
    q = db.query(models.InfraStat.section, models.InfraStat.id).filter(models.InfraStat.period == period) \
        if period else db.query(models.InfraStat.section, models.InfraStat.id)
    first_id = {}
    for section, sid in q.all():
        if section not in first_id or sid < first_id[section]:
            first_id[section] = sid
    return [s for s, _ in sorted(first_id.items(), key=lambda kv: kv[1])]


@admin_router.get("/csdl-ht/section")
def section_table(section: str, period: Optional[str] = None, db: Session = Depends(get_db),
                  _=Depends(require_module("csdl_ht", "view"))):
    """Bảng đúng như trong Excel gốc cho 1 nhóm: cột là từng chỉ tiêu (đúng
    thứ tự cột gốc), hàng là từng trung tâm rồi tới dòng "Tổng" cuối cùng."""
    period = period or _latest_period(db)
    q = db.query(models.InfraStat).filter(models.InfraStat.section == section)
    if period:
        q = q.filter(models.InfraStat.period == period)
    rows = q.order_by(models.InfraStat.id).all()

    columns, seen = [], set()
    by_center, center_order, tong_key = {}, [], None
    for r in rows:
        if r.center not in by_center:
            by_center[r.center] = {}
            if r.center.strip().casefold() in ("tổng", "tong"):
                tong_key = r.center
            else:
                center_order.append(r.center)
        by_center[r.center][r.label] = r.value if r.value is not None else r.value_text
        if r.label not in seen:
            seen.add(r.label)
            columns.append(r.label)

    center_order.sort()
    out_rows = [{"center": c, "values": [by_center[c].get(col) for col in columns]} for c in center_order]
    if tong_key:
        out_rows.append({"center": "Tổng", "values": [by_center[tong_key].get(col) for col in columns]})

    return {"period": period, "section": section, "columns": columns, "rows": out_rows}


@admin_router.get("/csdl-ht/export")
def export_csdl_csv(period: Optional[str] = None, center: Optional[str] = None,
                    db: Session = Depends(get_db), _=Depends(require_module("csdl_ht", "view"))):
    q = db.query(models.InfraStat)
    if period:
        q = q.filter(models.InfraStat.period == period)
    if center:
        q = q.filter(models.InfraStat.center == center)
    rows = q.order_by(models.InfraStat.id).all()
    header = ["Kỳ báo cáo", "Nhóm", "Trung tâm", "Chỉ tiêu", "Giá trị"]
    lines = [",".join(header)]
    for r in rows:
        vals = [r.period, r.section, r.center, r.label, r.value if r.value is not None else (r.value_text or "")]
        lines.append(",".join('"' + str(v).replace('"', '""') + '"' for v in vals))
    content = "﻿" + "\n".join(lines) + "\n"
    return PlainTextResponse(
        content, media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="csdl-ha-tang-{models.today()}.csv"'},
    )
