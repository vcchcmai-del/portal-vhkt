"""
Lịch trực toàn chi nhánh theo tháng — đội xe, chỉ huy, PVHKT và từng trung tâm.

Nhập bằng cách tải lên file Excel dạng "Tổng hợp" (đúng khuôn dạng file các
trung tâm đang tự làm và gửi qua lại): cột A là tên khối (PVHKT / Lái xe /
Chỉ Huy / tên trung tâm), cột B-F là mã NV/tên/SĐT/tổ nhóm/chức danh, từ cột G
là các cột ngày trong tháng, hai cột cuối là tên sự kiện hỗ trợ và địa điểm hỗ
trợ (nếu có). Dòng 2 (chỉ số 1) chứa ngày tháng dương lịch của từng cột ngày.

Nhập lại cho đúng khoảng ngày đã có sẽ thay thế toàn bộ dữ liệu cũ của khoảng
đó — không cộng dồn trùng lặp. Dữ liệu tháng khác không bị ảnh hưởng, vẫn tra
cứu lại được.
"""
import datetime as dt
import io
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from sqlalchemy.orm import Session

from . import models
from .auditlog import log_action
from .database import get_db
from .permissions import require_module

public_router = APIRouter(prefix="/api", tags=["Lịch trực toàn chi nhánh"])
admin_router = APIRouter(prefix="/api/admin", tags=["Quản trị lịch trực toàn chi nhánh"])

MAX_FILE_BYTES = 8 * 1024 * 1024


class RosterImportError(Exception):
    """Lỗi hiển thị được cho người dùng."""


def _excel_serial_to_date(serial: float) -> dt.date:
    return dt.date(1899, 12, 30) + dt.timedelta(days=round(serial))


def _cell_to_date(v) -> Optional[dt.date]:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and v > 0:
        return _excel_serial_to_date(v)
    return None


def _clean(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def parse_workbook(content: bytes) -> list:
    """Đọc file Excel, trả về danh sách dòng dạng dict khớp với cột của
    DutyRosterShift (chỉ những ngày có mã trực, bỏ qua ô rỗng)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RosterImportError("Máy chủ chưa cài thư viện đọc Excel (openpyxl).")

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise RosterImportError(f"Không đọc được file Excel: {e}")

    sheet_name = "Tổng hợp" if "Tổng hợp" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 3:
        raise RosterImportError("File thiếu dữ liệu — cần ít nhất dòng tiêu đề, dòng ngày và một dòng nhân sự.")

    header0 = rows[0] or []
    event_col = loc_col = -1
    for i, v in enumerate(header0):
        t = _clean(v)
        if t == "Tên sự kiện hỗ trợ":
            event_col = i
        elif t == "Địa điểm hỗ trợ":
            loc_col = i

    date_row = rows[1] or []
    date_cols = []
    for c in range(6, len(date_row)):
        d = _cell_to_date(date_row[c])
        if d is None:
            break
        date_cols.append((c, d))
    if not date_cols:
        raise RosterImportError("Không tìm thấy cột ngày (dòng 2, bắt đầu từ cột G).")

    out = []
    current_block = None
    for r in rows[2:]:
        if not r:
            continue
        col_a = _clean(r[0] if len(r) > 0 else "")
        col_b = _clean(r[1] if len(r) > 1 else "")
        if col_a:
            current_block = col_a
        if col_b and col_b != "Mã nhân viên":
            if not current_block:
                continue
            name = _clean(r[2] if len(r) > 2 else "")
            phone = _clean(r[3] if len(r) > 3 else "")
            group = _clean(r[4] if len(r) > 4 else "")
            title = _clean(r[5] if len(r) > 5 else "")
            event = _clean(r[event_col]) if event_col >= 0 and event_col < len(r) else ""
            location = _clean(r[loc_col]) if loc_col >= 0 and loc_col < len(r) else ""
            for col_idx, the_date in date_cols:
                code = _clean(r[col_idx] if col_idx < len(r) else "")
                if not code:
                    continue
                out.append({
                    "block": current_block, "duty_date": the_date, "employee_code": col_b,
                    "person_name": name, "phone": phone, "group_name": group, "title": title,
                    "shift_code": code, "event": event, "location": location,
                })
    if not out:
        raise RosterImportError("Không đọc được dòng nhân sự nào — kiểm tra lại đúng sheet 'Tổng hợp'.")
    return out


# ------------------------------------------------------------------- Tệp mẫu

SAMPLE_BLOCKS = [
    ("PVHKT", [
        ("NV001", "Nguyễn Văn A", "0900000001", "PVHKT", "Nhân viên Vận hành khai thác", ["TB", "", "TB"]),
    ]),
    ("Lái xe", [
        ("NV010", "Trần Văn B", "0900000010", "Đội xe", "Nhân viên lái xe", ["LX", "", "LX"]),
    ]),
    ("Chỉ Huy", [
        ("NV020", "Lê Văn C", "0900000020", "Trung tâm Thới Hòa", "Giám đốc kỹ thuật", ["CH", "CH", "CH"]),
    ]),
    ("Trung tâm Thới Hòa", [
        ("NV030", "Phạm Văn D", "0900000030", "Trung tâm Thới Hòa", "Nhân viên Kỹ thuật nhà trạm", ["KT", "", "KT"]),
    ]),
]


def build_roster_template_xlsx(month: Optional[str] = None) -> bytes:
    """
    Sinh tệp mẫu Excel đúng khuôn dạng "Tổng hợp" mà admin_import_roster() đọc
    được — kèm vài dòng ví dụ (PVHKT / Lái xe / Chỉ Huy / một trung tâm) và
    trang hướng dẫn, theo đúng phong cách tệp mẫu của mục Nhập dữ liệu hàng loạt.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if month:
        y, m = (int(x) for x in month.split("-"))
    else:
        today = dt.date.today()
        y, m = today.year, today.month
    start = dt.date(y, m, 1)
    end = dt.date(y + (m == 12), (m % 12) + 1, 1) - dt.timedelta(days=1)
    dates = [start + dt.timedelta(days=i) for i in range((end - start).days + 1)]

    fixed_cols = ["Đơn vị", "Mã nhân viên", "Tên nhân viên", "Số điện thoại", "Tổ nhóm", "Chức danh"]
    trailing_cols = ["Tên sự kiện hỗ trợ", "Địa điểm hỗ trợ"]
    total_cols = len(fixed_cols) + len(dates) + len(trailing_cols)

    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng hợp"

    do_bat_buoc = PatternFill("solid", fgColor="C8102E")
    xam_nhat = PatternFill("solid", fgColor="6B7280")
    chu_trang = Font(bold=True, color="FFFFFF", size=11)
    chu_vi_du = Font(italic=True, color="8A8A8A")

    for i, name in enumerate(fixed_cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill = do_bat_buoc; c.font = chu_trang
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = max(len(name) + 2, 14)
    for i, d in enumerate(dates, start=len(fixed_cols) + 1):
        c = ws.cell(row=1, column=i, value=d.strftime("%d/%m"))
        c.fill = xam_nhat; c.font = chu_trang
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = 7
        ws.cell(row=2, column=i, value=d).number_format = "dd/mm/yyyy"
    for j, name in enumerate(trailing_cols, start=len(fixed_cols) + len(dates) + 1):
        c = ws.cell(row=1, column=j, value=name)
        c.fill = xam_nhat; c.font = chu_trang
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = 22

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A3"

    row = 3
    for block_name, people in SAMPLE_BLOCKS:
        ws.cell(row=row, column=1, value=block_name).font = Font(bold=True, color="C8102E")
        row += 1
        for code, name, phone, group, title, shifts in people:
            ws.cell(row=row, column=2, value=code).font = chu_vi_du
            ws.cell(row=row, column=3, value=name).font = chu_vi_du
            ws.cell(row=row, column=4, value=phone).font = chu_vi_du
            ws.cell(row=row, column=5, value=group).font = chu_vi_du
            ws.cell(row=row, column=6, value=title).font = chu_vi_du
            for k, code_val in enumerate(shifts[:len(dates)]):
                if code_val:
                    ws.cell(row=row, column=len(fixed_cols) + 1 + k, value=code_val).font = chu_vi_du
            row += 1

    hd = wb.create_sheet("Huong dan")
    hd.column_dimensions["A"].width = 78
    hd["A1"] = "TỆP MẪU: Lịch trực toàn chi nhánh theo tháng"
    hd["A1"].font = Font(bold=True, size=14, color="C8102E")
    for i, line in enumerate([
        "",
        f"Trang 'Tổng hợp' đã điền sẵn đúng số ngày của tháng {m:02d}/{y} ở dòng 2 — không đổi số cột ngày.",
        "",
        "CẤU TRÚC BẢNG",
        "• 6 cột đầu (tô đỏ): Đơn vị, Mã nhân viên, Tên nhân viên, SĐT, Tổ nhóm, Chức danh.",
        "• Các cột ngày (tô xám, từ cột G): mỗi cột là một ngày trong tháng, dòng 2 đã có sẵn ngày.",
        "• 2 cột cuối (tô xám): Tên sự kiện hỗ trợ, Địa điểm hỗ trợ — chỉ điền khi người đó đi hỗ trợ nơi khác.",
        "",
        "CÁCH NHẬP",
        "1. Một dòng chỉ có chữ ở cột 'Đơn vị' (bỏ trống các cột khác) = bắt đầu một khối mới:",
        "   PVHKT, Lái xe, Chỉ Huy, hoặc tên trung tâm (giữ đúng chính tả xuyên suốt các tháng).",
        "2. Các dòng NGAY SAU đó, có Mã nhân viên, là người thuộc khối vừa mở.",
        "3. Ô ngày để trống = người đó KHÔNG trực hôm đó. Điền mã trực (VD: TB, CH, LX, KT,",
        "   hoặc mã riêng của đơn vị) vào đúng cột ngày người đó có trực.",
        "4. Xoá các dòng ví dụ (chữ nghiêng, xám) trước khi nhập dữ liệu thật.",
        "5. Lưu tệp rồi vào Quản trị > Nhập lịch trực toàn chi nhánh, chọn tệp này để tải lên.",
        "6. Nhập lại đúng khoảng ngày đã có sẽ THAY THẾ dữ liệu cũ của khoảng đó, không tạo trùng.",
        "   Dữ liệu của tháng khác không bị ảnh hưởng.",
        "7. Giữ nguyên tên trang tính là 'Tổng hợp' — hệ thống tìm đúng tên này để đọc dữ liệu.",
    ], start=2):
        hd.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@admin_router.get("/duty-roster/template")
def admin_roster_template(month: Optional[str] = None, _=Depends(require_module("duty_roster", "view"))):
    """Tải tệp mẫu Excel đúng khuôn dạng để nhập lịch trực."""
    from fastapi.responses import Response
    try:
        content = build_roster_template_xlsx(month)
    except ImportError:
        raise HTTPException(500, "Máy chủ chưa cài thư viện tạo tệp Excel (openpyxl).")
    except (ValueError, TypeError):
        raise HTTPException(400, "Định dạng tháng không hợp lệ, cần dạng YYYY-MM.")
    label = month or dt.date.today().strftime("%Y-%m")
    return Response(
        content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="mau-lich-truc-{label}.xlsx"'},
    )


# --------------------------------------------------------------------- Nhập

@admin_router.post("/duty-roster/import")
async def admin_import_roster(file: UploadFile = File(...), db: Session = Depends(get_db),
                              user=Depends(require_module("duty_roster", "create")), request: Request = None):
    content = await file.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(400, f"Tệp vượt quá {MAX_FILE_BYTES // (1024*1024)} MB.")
    try:
        parsed = parse_workbook(content)
    except RosterImportError as e:
        raise HTTPException(400, str(e))

    dates = sorted({row["duty_date"] for row in parsed})
    blocks = sorted({row["block"] for row in parsed})
    people = {(row["block"], row["employee_code"], row["person_name"]) for row in parsed}

    db.query(models.DutyRosterShift).filter(
        models.DutyRosterShift.duty_date >= dates[0],
        models.DutyRosterShift.duty_date <= dates[-1],
    ).delete()
    db.bulk_insert_mappings(models.DutyRosterShift, parsed)
    db.commit()

    label = f"{dates[0].isoformat()} → {dates[-1].isoformat()}"
    log_action(db, user, "import", "duty_roster", None, label,
              detail=f"{len(people)} người, {len(blocks)} khối, {len(dates)} ngày", request=request)
    return {"rows": len(parsed), "people": len(people), "blocks": blocks,
            "date_from": dates[0], "date_to": dates[-1]}


@admin_router.delete("/duty-roster/months/{month}")
def admin_delete_roster_month(month: str, db: Session = Depends(get_db),
                              user=Depends(require_module("duty_roster", "delete")), request: Request = None):
    """Xoá toàn bộ dữ liệu của một tháng, ví dụ month='2026-08'."""
    try:
        y, m = (int(x) for x in month.split("-"))
        start = dt.date(y, m, 1)
        end = dt.date(y + (m == 12), (m % 12) + 1, 1) - dt.timedelta(days=1)
    except (ValueError, TypeError):
        raise HTTPException(400, "Định dạng tháng không hợp lệ, cần dạng YYYY-MM.")
    count = db.query(models.DutyRosterShift).filter(
        models.DutyRosterShift.duty_date >= start, models.DutyRosterShift.duty_date <= end,
    ).delete()
    db.commit()
    log_action(db, user, "delete", "duty_roster", None, month, detail=f"{count} dòng", request=request)
    return {"deleted_rows": count}


# --------------------------------------------------------------------- Tra cứu

@public_router.get("/duty-roster/months")
def list_roster_months(db: Session = Depends(get_db)):
    rows = db.query(models.DutyRosterShift.duty_date).distinct().all()
    months = sorted({d[0].strftime("%Y-%m") for d in rows}, reverse=True)
    return months


def _build_roster(db: Session, start: dt.date, end: dt.date) -> dict:
    rows = (
        db.query(models.DutyRosterShift)
        .filter(models.DutyRosterShift.duty_date >= start, models.DutyRosterShift.duty_date <= end)
        .order_by(models.DutyRosterShift.duty_date)
        .all()
    )
    if not rows:
        return {"dates": [], "blocks": []}

    dates = sorted({r.duty_date for r in rows})
    date_index = {d: i for i, d in enumerate(dates)}

    block_order = []
    block_people = {}   # block -> { (code,name) -> person dict }
    for r in rows:
        if r.block not in block_people:
            block_people[r.block] = {}
            block_order.append(r.block)
        key = (r.employee_code, r.person_name)
        people = block_people[r.block]
        if key not in people:
            people[key] = {
                "code": r.employee_code, "name": r.person_name, "phone": r.phone,
                "group": r.group_name, "title": r.title,
                "shifts": ["" for _ in dates], "event": "", "location": "",
            }
        p = people[key]
        p["shifts"][date_index[r.duty_date]] = r.shift_code or ""
        if r.event:
            p["event"] = r.event
        if r.location:
            p["location"] = r.location

    blocks = [{"name": b, "people": list(block_people[b].values())} for b in block_order]
    return {"dates": [d.isoformat() for d in dates], "blocks": blocks}


@public_router.get("/duty-roster")
def get_roster(month: Optional[str] = None, db: Session = Depends(get_db)):
    """Toàn bộ lịch trực của một tháng. Không truyền `month` thì lấy tháng
    gần nhất đang có dữ liệu."""
    if not month:
        latest = list_roster_months(db)
        if not latest:
            return {"dates": [], "blocks": [], "month": None}
        month = latest[0]
    try:
        y, m = (int(x) for x in month.split("-"))
        start = dt.date(y, m, 1)
        end = dt.date(y + (m == 12), (m % 12) + 1, 1) - dt.timedelta(days=1)
    except (ValueError, TypeError):
        raise HTTPException(400, "Định dạng tháng không hợp lệ, cần dạng YYYY-MM.")
    data = _build_roster(db, start, end)
    data["month"] = month
    return data
